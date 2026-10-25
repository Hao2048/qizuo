# coding: utf-8
"""Local acceptance fixtures only; requires python-docx/openpyxl/python-pptx/reportlab/Pillow and ffmpeg."""
from pathlib import Path
import subprocess
import sys
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw
p = Path(sys.argv[1] if len(sys.argv) > 1 else '/tmp/pih-artifact-fixtures')
p.mkdir(parents=True, exist_ok=True)
doc = Document(); doc.add_heading('Artifact delivery acceptance', 0); doc.add_paragraph('Document preview is readable.'); doc.save(p / 'report.docx')
w = Workbook(); w.active.append(['Metric', 'Value']); w.active.append(['Pass', 42]); w.save(p / 'metrics.xlsx')
r = Presentation(); slide = r.slides.add_slide(r.slide_layouts[0]); slide.shapes.title.text = 'Delivery acceptance'; r.save(p / 'slides.pptx')
c = canvas.Canvas(str(p / 'report.pdf')); c.drawString(72, 740, 'PDF DELIVERY ACCEPTANCE'); c.save()
im = Image.new('RGB', (640, 360), '#eef2ff'); ImageDraw.Draw(im).rectangle((80, 100, 450, 250), fill='#4f46e5'); im.save(p / 'chart.png')
(p / 'chart.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="640" height="360"><rect width="640" height="360" fill="#eef2ff"/><rect x="80" y="100" width="400" height="150" fill="#4f46e5"/><text x="80" y="70" font-size="24">SVG DELIVERY CHART</text></svg>')
(p / 'PRD.md').write_text('# Delivery PRD\n\n## Goal\n\nReadable artifact delivery.\n\n| Check | Result |\n|---|---|\n| Document | Pass |\n')
(p / 'demo.ts').write_text('export const result: string = "DELIVERY CODE";\n')
(p / 'index.html').write_text('<!doctype html><meta charset="utf-8"><title>Artifact acceptance</title><h1>DELIVERY WEBSITE</h1><img src="chart.svg" width="320"><p>Static website works.</p>')
subprocess.run(['ffmpeg', '-hide_banner', '-loglevel', 'error', '-f', 'lavfi', '-i', 'testsrc2=size=640x360:rate=24', '-f', 'lavfi', '-i', 'sine=frequency=440:sample_rate=44100', '-t', '3', '-c:v', 'libx264', '-pix_fmt', 'yuv420p', '-c:a', 'aac', '-movflags', '+faststart', '-y', str(p / 'clip.mp4')], check=True)
